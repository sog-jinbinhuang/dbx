"""
order_audit.py
====================
Reads from FCT_GLOBAL_ORDER_AUDIT (dbt model) and renders an Excel
workbook with one tab per SOURCE_DATABASE, listing who entered each
order, the customer, order number, revenue, and quantity.

Tabs produced
-------------
  SUMMARY        — order count, total revenue, total qty per database
  <DATABASE>     — one tab per database: Created By / Customer /
                    Order Num / Revenue / Quantity (kg), sorted by
                    revenue descending, with a grand total row

Run standalone:  python order_audit.py
Used by:         reports/utils/finance_emailer.py (legacy) and
                 report_registry.py / send_reports.py (current)
"""

from __future__ import annotations

import io
import os
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from databricks import sql as databricks_sql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# ═══════════════════════════════════════════════════════════════════════════════
# CONFIG
# ═══════════════════════════════════════════════════════════════════════════════

_DIR = os.path.dirname(os.path.abspath(__file__))

DBX_SERVER_HOSTNAME = os.environ["DATABRICKS_SERVER_HOSTNAME"]
DBX_HTTP_PATH       = os.environ["DATABRICKS_HTTP_PATH"]
DBX_ACCESS_TOKEN    = os.environ["DATABRICKS_TOKEN"]
DBX_CATALOG         = "dev"
DBX_SCHEMA          = "gold_finance"   # matches AR/AP -- confirm FCT_GLOBAL_ORDER_AUDIT actually lands here

MODEL_NAME = "FCT_GLOBAL_ORDER_AUDIT"

TODAY        = date.today()
_DATE_SUFFIX = TODAY.strftime("%m%d%Y")
OUTPUT_FILE  = os.path.join(_DIR, f"order_audit_{_DATE_SUFFIX}.xlsx")

NUM_QTY = "#,##0.00;-#,##0.00"
NUM_USD = "#,##0.00;-#,##0.00"

# ── emailer compatibility ──────────────────────────────────────────────────────

@dataclass
class _ReportConfig:
    output_file: str | Path | io.IOBase

ORDER_AUDIT_CONFIG = _ReportConfig(output_file=OUTPUT_FILE)


# ═══════════════════════════════════════════════════════════════════════════════
# STYLESHEET  (lighter slate-blue theme, matching rpt_backlog.py / rpt_backlog_china.py)
# ═══════════════════════════════════════════════════════════════════════════════

class S:
    TITLE_BG = "34495E"; TITLE_FG = "FFFFFF"
    GRP_BG   = "5B84B1"; GRP_FG   = "FFFFFF"
    COL_BG   = "DCE6F1"; COL_FG   = "000000"
    GRAND_BG = "34495E"; GRAND_FG = "FFFFFF"
    ALT_BG   = "F7F9FC"; EVEN_BG  = "FFFFFF"
    SEP      = Side(style="thin",   color="8FA3BF")
    OUTER    = Side(style="medium", color="6C7A89")
    THIN     = Side(style="thin",   color="B0BAC5")
    THIN_G   = Side(style="thin",   color="E1E5EA")

    @staticmethod
    def fill(h: str) -> PatternFill: return PatternFill("solid", fgColor=h)
    @staticmethod
    def font(bold=False, color="000000", size=9, italic=False) -> Font:
        return Font(name="Arial", bold=bold, color=color, size=size, italic=italic)
    @staticmethod
    def align(h="right", v="center", wrap=False) -> Alignment:
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    @classmethod
    def right_sep(cls) -> Border: return Border(right=cls.SEP)

    @classmethod
    def apply_outer_border(cls, ws: Worksheet,
                           r1: int, r2: int, c1: int, c2: int) -> None:
        for r in range(r1, r2 + 1):
            lc = ws.cell(row=r, column=c1)
            rc = ws.cell(row=r, column=c2)
            lc.border = Border(left=cls.OUTER,  top=lc.border.top,
                               bottom=lc.border.bottom, right=lc.border.right)
            rc.border = Border(right=cls.OUTER, top=rc.border.top,
                               bottom=rc.border.bottom, left=rc.border.left)


# ═══════════════════════════════════════════════════════════════════════════════
# LOADER
# ═══════════════════════════════════════════════════════════════════════════════

class Loader:
    def __init__(self) -> None:
        print(f"Connecting to Databricks ({DBX_SERVER_HOSTNAME}) ...")
        self._conn = databricks_sql.connect(
            server_hostname = DBX_SERVER_HOSTNAME,
            http_path       = DBX_HTTP_PATH,
            access_token    = DBX_ACCESS_TOKEN,
            catalog         = DBX_CATALOG,
            schema          = DBX_SCHEMA,
        )
        print("  Connected.")

    def close(self) -> None:
        self._conn.close()
        print("  Databricks connection closed.")

    def load(self) -> pd.DataFrame:
        # "DATABASE" is a reserved word -- Databricks SQL uses backticks for
        # quoted identifiers (not double quotes, which Spark treats as a
        # string literal unless ANSI quoted-identifier mode is on).
        query = f"""
            SELECT
                `DATABASE`  AS SOURCE_DATABASE,
                CREATED_BY,
                CUST_NAME,
                ORDER_NUM,
                REVENUE,
                QTY_IN_KG
            FROM {MODEL_NAME}
        """
        print(f"  Loading {MODEL_NAME} ...")
        cur = self._conn.cursor()
        cur.execute(query)
        cols = [d[0].upper() for d in cur.description]
        rows = cur.fetchall()
        cur.close()

        df = pd.DataFrame(rows, columns=cols)
        print(f"    {len(df):,} rows returned")

        for col in ["SOURCE_DATABASE", "CREATED_BY", "CUST_NAME", "ORDER_NUM"]:
            df[col] = (df[col].astype(str).str.strip()
                       .replace({"nan": "", "None": ""}).fillna(""))
        for col in ["REVENUE", "QTY_IN_KG"]:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

        df = df.sort_values(
            ["SOURCE_DATABASE", "REVENUE"], ascending=[True, False]
        ).reset_index(drop=True)
        return df


# ═══════════════════════════════════════════════════════════════════════════════
# REPORT BUILDER
# ═══════════════════════════════════════════════════════════════════════════════

class OrderAuditReportBuilder:

    def __init__(self, wb: Workbook, df: pd.DataFrame) -> None:
        self.wb = wb
        self.df = df

    # ── low-level helpers ─────────────────────────────────────────────────────

    def _sc(self, ws, row, col, value=None, *, bold=False, fg="000000",
            bg="FFFFFF", fmt=None, halign="right", border=None,
            italic=False, wrap=False):
        c = ws.cell(row=row, column=col, value=value)
        c.font      = S.font(bold=bold, color=fg, italic=italic)
        c.fill      = bg if isinstance(bg, PatternFill) else S.fill(bg)
        c.alignment = S.align(halign, wrap=wrap)
        if fmt    is not None: c.number_format = fmt
        if border is not None: c.border        = border
        return c

    def _title_banner(self, ws: Worksheet, total_cols: int,
                      title: str, subtitle: str) -> None:
        for row, text, bold, size, color in [
            (1, title,    True,  11, S.TITLE_FG),
            (2, subtitle, False,  9, S.COL_BG),
        ]:
            ws.merge_cells(start_row=row, start_column=1,
                           end_row=row,   end_column=total_cols)
            c = ws.cell(row=row, column=1, value=text)
            c.font      = S.font(bold=bold, color=color, size=size)
            c.fill      = S.fill(S.TITLE_BG)
            c.alignment = S.align("center")
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 14

    def _databases_sorted(self) -> list[str]:
        return sorted(self.df["SOURCE_DATABASE"].unique())

    # ── SUMMARY tab ───────────────────────────────────────────────────────────

    def build_summary_tab(self) -> None:
        ws = self.wb.create_sheet("SUMMARY", index=0)
        TOTAL_COLS = 4

        self._title_banner(
            ws, TOTAL_COLS,
            f"Global Order Audit  |  {TODAY.strftime('%B %d, %Y')}",
            "Source: FCT_GLOBAL_ORDER_AUDIT",
        )

        headers = ["Database", "Orders", "Revenue (USD)", "Quantity (kg)"]
        widths  = [18, 12, 18, 18]
        for ci, (hdr, width) in enumerate(zip(headers, widths), start=1):
            c = ws.cell(row=3, column=ci, value=hdr)
            c.font      = S.font(bold=True)
            c.fill      = S.fill(S.COL_BG)
            c.alignment = S.align("center", wrap=True)
            c.border    = Border(left=S.THIN_G, right=S.THIN_G,
                                 top=S.OUTER,   bottom=S.OUTER)
            ws.column_dimensions[get_column_letter(ci)].width = width
        ws.row_dimensions[3].height = 18

        databases = self._databases_sorted()
        cur_row   = 4
        grand_orders, grand_rev, grand_qty = 0, 0.0, 0.0

        for di, db in enumerate(databases):
            sub     = self.df[self.df["SOURCE_DATABASE"] == db]
            orders  = len(sub)
            revenue = float(sub["REVENUE"].sum())
            qty     = float(sub["QTY_IN_KG"].sum())
            row_bg  = S.ALT_BG if di % 2 == 1 else S.EVEN_BG

            self._sc(ws, cur_row, 1, db, bold=True, bg=row_bg, halign="left",
                     border=Border(bottom=S.THIN_G))
            self._sc(ws, cur_row, 2, orders, bg=row_bg, fmt="#,##0",
                     halign="center", border=Border(bottom=S.THIN_G))
            self._sc(ws, cur_row, 3, revenue, bg=row_bg, fmt=NUM_USD,
                     border=Border(bottom=S.THIN_G))
            self._sc(ws, cur_row, 4, qty, bg=row_bg, fmt=NUM_QTY,
                     border=Border(bottom=S.THIN_G))

            grand_orders += orders
            grand_rev    += revenue
            grand_qty    += qty
            ws.row_dimensions[cur_row].height = 15
            cur_row += 1

        self._sc(ws, cur_row, 1, "Grand Total", bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, halign="left")
        self._sc(ws, cur_row, 2, grand_orders, bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, fmt="#,##0", halign="center")
        self._sc(ws, cur_row, 3, grand_rev, bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, fmt=NUM_USD)
        self._sc(ws, cur_row, 4, grand_qty, bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, fmt=NUM_QTY)
        ws.row_dimensions[cur_row].height = 16

        S.apply_outer_border(ws, 1, cur_row, 1, TOTAL_COLS)
        ws.freeze_panes = "A4"
        print(f"  Summary: {len(databases)} databases | {grand_orders:,} orders | "
              f"${grand_rev:,.0f} revenue")

    # ── per-database DETAIL tab ──────────────────────────────────────────────

    def build_detail_tab(self, db_df: pd.DataFrame, sheet_name: str, db: str) -> None:
        ws = self.wb.create_sheet(sheet_name)
        cols = [
            ("Created By",     22, "CREATED_BY",  "@",      "left"),
            ("Customer",       30, "CUST_NAME",   "@",      "left"),
            ("Order Num",      16, "ORDER_NUM",   "@",      "center"),
            ("Revenue (USD)",  16, "REVENUE",     NUM_USD,  "right"),
            ("Quantity (kg)",  16, "QTY_IN_KG",   NUM_QTY,  "right"),
        ]
        N = len(cols)

        self._title_banner(
            ws, N,
            f"Global Order Audit  |  {db}  |  {TODAY.strftime('%B %d, %Y')}",
            f"{len(db_df):,} orders  |  Source: FCT_GLOBAL_ORDER_AUDIT",
        )
        ws.row_dimensions[3].height = 20

        for ci, (hdr, width, _, _, _) in enumerate(cols, start=1):
            c = ws.cell(row=3, column=ci, value=hdr)
            c.font      = S.font(bold=True)
            c.fill      = S.fill(S.COL_BG)
            c.alignment = S.align("center", wrap=True)
            c.border    = Border(left=S.THIN_G, right=S.THIN_G,
                                 top=S.OUTER,   bottom=S.OUTER)
            ws.column_dimensions[get_column_letter(ci)].width = width

        cur_row = 4
        rev_total, qty_total = 0.0, 0.0
        for row_idx, (_, r) in enumerate(db_df.iterrows()):
            row_bg = S.ALT_BG if row_idx % 2 == 1 else S.EVEN_BG
            ws.row_dimensions[cur_row].height = 15
            for ci, (_, _, key, fmt, halign) in enumerate(cols, start=1):
                val = r[key]
                c = ws.cell(row=cur_row, column=ci, value=val)
                c.font          = S.font(bold=False)
                c.fill          = S.fill(row_bg)
                c.alignment     = S.align(halign)
                c.number_format = fmt
                c.border        = Border(left=S.THIN_G, right=S.THIN_G, bottom=S.THIN_G)
            rev_total += float(r["REVENUE"])
            qty_total += float(r["QTY_IN_KG"])
            cur_row += 1

        # grand total row
        self._sc(ws, cur_row, 1, "GRAND TOTAL", bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, halign="left")
        self._sc(ws, cur_row, 2, None, bg=S.GRAND_BG)
        self._sc(ws, cur_row, 3, None, bg=S.GRAND_BG)
        self._sc(ws, cur_row, 4, rev_total, bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, fmt=NUM_USD)
        self._sc(ws, cur_row, 5, qty_total, bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, fmt=NUM_QTY)
        ws.row_dimensions[cur_row].height = 16

        S.apply_outer_border(ws, 1, cur_row, 1, N)
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(N)}3"
        print(f"  {sheet_name}: {len(db_df):,} orders | ${rev_total:,.0f} revenue")

    # ── orchestrate ───────────────────────────────────────────────────────────

    def build(self) -> None:
        self.build_summary_tab()
        for db in self._databases_sorted():
            safe = (db[:31] or "UNKNOWN")
            safe = (safe.replace("/", "-").replace("\\", "-")
                    .replace("*", "").replace("?", "")
                    .replace("[", "").replace("]", "")
                    .replace(":", "").upper())
            db_df = self.df[self.df["SOURCE_DATABASE"] == db]
            self.build_detail_tab(db_df, safe, db)


# ═══════════════════════════════════════════════════════════════════════════════
# EMAILER ENTRY POINT (legacy -- kept for finance_emailer.py compatibility)
# ═══════════════════════════════════════════════════════════════════════════════

def build_report(df: pd.DataFrame, _orders_df: Any, config: _ReportConfig) -> None:
    """Called by finance_emailer.py. _orders_df is unused (this report has no
    separate orders dataframe -- df IS the order data)."""
    wb = Workbook()
    wb.remove(wb.active)
    builder = OrderAuditReportBuilder(wb, df)
    builder.build()
    wb.save(config.output_file)


# ═══════════════════════════════════════════════════════════════════════════════
# STANDARD ENTRY POINT (used by report_registry.py / send_reports.py)
# ═══════════════════════════════════════════════════════════════════════════════

def build_attachments() -> list[tuple[str, bytes]]:
    """Loads data, builds the workbook in memory, returns [(filename, bytes)]."""
    loader = Loader()
    try:
        df = loader.load()
    finally:
        loader.close()

    wb = Workbook()
    wb.remove(wb.active)
    builder = OrderAuditReportBuilder(wb, df)
    builder.build()

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return [(f"order_audit_{_DATE_SUFFIX}.xlsx", buf.getvalue())]


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════

def main() -> None:
    loader = Loader()
    try:
        df = loader.load()
    finally:
        loader.close()

    wb = Workbook()
    wb.remove(wb.active)

    builder = OrderAuditReportBuilder(wb, df)
    builder.build()

    wb.save(ORDER_AUDIT_CONFIG.output_file)
    print(f"\nSaved -> {ORDER_AUDIT_CONFIG.output_file}")
    print(f"Tabs:  {[ws.title for ws in wb.worksheets]}")


if __name__ == "__main__":
    main()
