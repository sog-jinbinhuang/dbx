"""
ar_report.py
============
Reads from FCT_GLOBAL_OPEN_AR and renders an Excel AR aging workbook.

Tabs produced
-------------
  SUMMARY        — one row per database, columns by aging bucket (USD), plus
                   a second table below showing each bucket as % of total open
  <DATABASE>     — one tab per source database, sorted days overdue desc→asc
  EXCHANGE RATES — spot rates used for USD conversion

Run standalone:  python ar_report.py
Used by:         reports/utils/finance_emailer.py

Requires env vars:
  DATABRICKS_SERVER_HOSTNAME   e.g. adb-1234567890123456.7.azuredatabricks.net
  DATABRICKS_HTTP_PATH         e.g. /sql/1.0/warehouses/abc123def456
  DATABRICKS_TOKEN             a personal access token
"""

from __future__ import annotations

import io
import os
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from databricks import sql as databricks_sql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# ═══════════════════════════════════════════════════════════════════════════════
# CONFIG
# ═══════════════════════════════════════════════════════════════════════════════

_DIR        = os.path.dirname(os.path.abspath(__file__))

DBX_SERVER_HOSTNAME = os.environ["DATABRICKS_SERVER_HOSTNAME"]
DBX_HTTP_PATH       = os.environ["DATABRICKS_HTTP_PATH"]
DBX_ACCESS_TOKEN    = os.environ["DATABRICKS_TOKEN"]

DBX_CATALOG = "dev"
DBX_SCHEMA  = "gold_finance"

TODAY        = date.today()
_DATE_SUFFIX = TODAY.strftime("%m%d%Y")
OUTPUT_FILE  = os.path.join(_DIR, f"ar_report_{_DATE_SUFFIX}.xlsx")

AGING_BUCKETS = [
    "Due > 30 D",
    "Due < 30 D",
    "PD 0 to 30 D",
    "PD 31 to 60 D",
    "PD 61 to 90 D",
    "PD > 90 D",
]

AGING_LABELS: dict[str, str] = {
    "Due > 30 D":    "Not Due\n(> 30 Days)",
    "Due < 30 D":    "Not Due\n(<= 30 Days)",
    "PD 0 to 30 D":  "Past Due\n(0 – 30 Days)",
    "PD 31 to 60 D": "Past Due\n(31 – 60 Days)",
    "PD 61 to 90 D": "Past Due\n(61 – 90 Days)",
    "PD > 90 D":     "Past Due\n(> 90 Days)",
}

DETAIL_COLS: list[tuple[str, int, str, str, str]] = [
    ("Cust Code",             10, "CUST_CODE",                  "@",                  "center"),
    ("Customer",              30, "CUST_NAME",                   "@",                  "left"),
    ("Sales Agent",           16, "DEFAULT_SALES_AGENT_2",     "@",                  "center"),
    ("Ref Type",               9, "REF_TYPE",                    "@",                  "center"),
    ("Ref #",                 12, "REF_NUM",                     "@",                  "center"),
    ("Ref Date",              12, "REF_DATE",                    "MM/DD/YYYY",          "center"),
    ("Due Date",              12, "DUE_DATE",                    "MM/DD/YYYY",          "center"),
    ("Days Overdue",          12, "DAYS_OVERDUE",                "#,##0;-#,##0",        "center"),
    ("Aging",                 18, "OPEN_GROUP_LABEL",            "@",                  "center"),
    ("Orig Amount\n(Local)",  16, "AMOUNT_LOCAL_CURRENCY",       "#,##0.00;-#,##0.00", "right"),
    ("Open Amount\n(Local)",  16, "OPEN_AMOUNT_LOCAL_CURRENCY",  "#,##0.00;-#,##0.00", "right"),
    ("Open Amount\n(USD)",    16, "OPEN_AMOUNT_IN_USD",          "#,##0.00;-#,##0.00", "right"),
]

NUM_USD = "#,##0;-#,##0;\"-\""

BUCKET_STYLE: dict[str, tuple[str, bool]] = {
    "Due > 30 D":    ("808080", False),
    "Due < 30 D":    ("595959", False),
    "PD 0 to 30 D":  ("BF8F00", False),
    "PD 31 to 60 D": ("C55A11", False),
    "PD 61 to 90 D": ("C00000", True),
    "PD > 90 D":     ("7B0000", True),
}

# ── emailer compatibility ──────────────────────────────────────────────────────

@dataclass
class _ReportConfig:
    output_file: str | Path | io.IOBase

AR_CONFIG = _ReportConfig(output_file=OUTPUT_FILE)


# ═══════════════════════════════════════════════════════════════════════════════
# STYLESHEET
# ═══════════════════════════════════════════════════════════════════════════════

class S:
    TITLE_BG = "34495E"; TITLE_FG = "FFFFFF"
    GRP_BG   = "5B84B1"; GRP_FG   = "FFFFFF"
    COL_BG   = "DCE6F1"; COL_FG   = "000000"
    GRAND_BG = "34495E"; GRAND_FG = "FFFFFF"
    ALT_BG   = "F7F9FC"; EVEN_BG  = "FFFFFF"
    SEP      = Side(style="thin",   color="8FA3BF")
    OUTER    = Side(style="medium", color="6C7A89")
    THIN     = Side(style="thin",   color="B0BAC5")
    THIN_G   = Side(style="thin",   color="E1E5EA")

    @staticmethod
    def fill(h: str) -> PatternFill: return PatternFill("solid", fgColor=h)
    @staticmethod
    def font(bold=False, color="000000", size=9, italic=False) -> Font:
        return Font(name="Arial", bold=bold, color=color, size=size, italic=italic)
    @staticmethod
    def align(h="right", v="center", wrap=False) -> Alignment:
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    @classmethod
    def right_sep(cls) -> Border: return Border(right=cls.SEP)

    @classmethod
    def apply_outer_border(cls, ws: Worksheet,
                           r1: int, r2: int, c1: int, c2: int) -> None:
        for r in range(r1, r2 + 1):
            lc = ws.cell(row=r, column=c1)
            rc = ws.cell(row=r, column=c2)
            lc.border = Border(left=cls.OUTER,  top=lc.border.top,
                               bottom=lc.border.bottom, right=lc.border.right)
            rc.border = Border(right=cls.OUTER, top=rc.border.top,
                               bottom=rc.border.bottom, left=rc.border.left)


# ═══════════════════════════════════════════════════════════════════════════════
# LOADER
# ═══════════════════════════════════════════════════════════════════════════════

class Loader:
    def __init__(self) -> None:
        print(f"Connecting to Databricks ({DBX_SERVER_HOSTNAME}) ...")
        self._conn = databricks_sql.connect(
            server_hostname = DBX_SERVER_HOSTNAME,
            http_path        = DBX_HTTP_PATH,
            access_token     = DBX_ACCESS_TOKEN,
            catalog          = DBX_CATALOG,
            schema           = DBX_SCHEMA,
        )
        print("  Connected.")

    def close(self) -> None:
        self._conn.close()
        print("  Databricks connection closed.")

    def load(self) -> pd.DataFrame:
        print("  Loading FCT_GLOBAL_OPEN_AR ...")
        cur = self._conn.cursor()
        cur.execute("SELECT * FROM FCT_GLOBAL_OPEN_AR")
        cols = [d[0].upper() for d in cur.description]
        rows = cur.fetchall()
        cur.close()
        df = pd.DataFrame(rows, columns=cols)
        print(f"    {len(df):,} rows, {len(df.columns)} columns")

        numeric_cols = ["AMT", "OP_AMT", "ENT_CURRENCY_CONVERSION_RATE",
                        "AMOUNT_LOCAL_CURRENCY", "OPEN_AMOUNT_LOCAL_CURRENCY",
                        "LATEST_X_RATE", "AMOUNT_IN_USD", "OPEN_AMOUNT_IN_USD",
                        "DAYS_OVERDUE"]
        for col in numeric_cols:
            if col in df.columns:
                df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

        for col in ["REF_DATE", "DUE_DATE"]:
            if col in df.columns:
                df[col] = pd.to_datetime(df[col], errors="coerce")

        str_cols = ["DATABASE", "DIVISION", "COMPANY", "CUST_CODE", "CUST_NAME", "DEFAULT_SALES_AGENT_2",
                    "REF_TYPE", "REF_NUM", "LOCAL_CURRENCY_LABEL",
                    "ENTERED_CURRENCY", "OPEN_GROUP", "GL_ACCT"]
        for col in str_cols:
            if col in df.columns:
                df[col] = (df[col].astype(str).str.strip()
                           .replace({"nan": "", "None": ""}).fillna(""))

        df["OPEN_GROUP_LABEL"] = df["OPEN_GROUP"].map(
            {k: v.replace("\n", " ") for k, v in AGING_LABELS.items()}
        ).fillna(df["OPEN_GROUP"])

        df = df.sort_values("DAYS_OVERDUE", ascending=False).reset_index(drop=True)
        return df

    def load_new_orders(self):
        """Stub — required by finance_emailer.py interface."""
        return None


# ═══════════════════════════════════════════════════════════════════════════════
# REPORT BUILDER
# ═══════════════════════════════════════════════════════════════════════════════

class ARReportBuilder:

    def __init__(self, wb: Workbook, df: pd.DataFrame) -> None:
        self.wb = wb
        self.df = df

    # ── low-level helpers ─────────────────────────────────────────────────────

    def _sc(self, ws: Worksheet, row: int, col: int, value: Any = None, *,
            bold=False, fg="000000", bg: str | PatternFill = "FFFFFF",
            fmt: str | None = None, halign="right", border=None,
            italic=False, wrap=False):
        c = ws.cell(row=row, column=col, value=value)
        c.font      = S.font(bold=bold, color=fg, italic=italic)
        c.fill      = bg if isinstance(bg, PatternFill) else S.fill(bg)
        c.alignment = S.align(halign, wrap=wrap)
        if fmt    is not None: c.number_format = fmt
        if border is not None: c.border        = border
        return c

    def _title_banner(self, ws: Worksheet, total_cols: int,
                      title: str, subtitle: str) -> None:
        for row, text, bold, size, color in [
            (1, title,    True,  11, S.TITLE_FG),
            (2, subtitle, False,  9, "BDD7EE"),
        ]:
            ws.merge_cells(start_row=row, start_column=1,
                           end_row=row,   end_column=total_cols)
            c = ws.cell(row=row, column=1, value=text)
            c.font      = S.font(bold=bold, color=color, size=size)
            c.fill      = S.fill(S.TITLE_BG)
            c.alignment = S.align("center")
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 14

    def _databases_sorted(self) -> list[str]:
        return sorted(self.df["DATABASE"].unique())

    # ── SUMMARY tab ───────────────────────────────────────────────────────────

    def build_summary_tab(self) -> None:
        ws = self.wb.create_sheet("SUMMARY", index=0)

        N_BUCKETS        = len(AGING_BUCKETS)
        COL_DB           = 1
        COL_BUCKET_START = 2
        COL_TOTAL        = COL_BUCKET_START + N_BUCKETS
        COL_PCT          = COL_TOTAL + 1
        TOTAL_COLS       = COL_PCT

        self._title_banner(
            ws, TOTAL_COLS,
            f"Mafco Open AR Summary  |  {TODAY.strftime('%B %d, %Y')}",
            "Open amounts in USD  |  Excludes intercompany balances",
        )

        # ── USD TABLE ─────────────────────────────────────────────────────────

        ws.merge_cells(start_row=3, start_column=COL_BUCKET_START,
                       end_row=3,   end_column=COL_TOTAL - 1)
        c3 = ws.cell(row=3, column=COL_BUCKET_START,
                     value="Aging Buckets  (Open Amount in USD)")
        c3.font      = S.font(bold=True, color=S.GRP_FG)
        c3.fill      = S.fill(S.GRP_BG)
        c3.alignment = S.align("center")
        for col in range(COL_BUCKET_START, COL_TOTAL):
            ws.cell(row=3, column=col).fill = S.fill(S.GRP_BG)
        for col in [COL_DB, COL_TOTAL, COL_PCT]:
            ws.cell(row=3, column=col).fill = S.fill(S.TITLE_BG)
        ws.row_dimensions[3].height = 14

        self._sc(ws, 4, COL_DB, "Database", bold=True, bg=S.COL_BG,
                 halign="left", border=S.right_sep())
        self._sc(ws, 4, COL_TOTAL, "Total Open\n(USD)", bold=True, bg=S.COL_BG,
                 halign="right", border=S.right_sep(), wrap=True)
        self._sc(ws, 4, COL_PCT, "% Past Due", bold=True, bg=S.COL_BG,
                 halign="center", border=S.right_sep(), wrap=True)
        for bi, bucket in enumerate(AGING_BUCKETS):
            col   = COL_BUCKET_START + bi
            label = AGING_LABELS[bucket]
            c     = ws.cell(row=4, column=col, value=label)
            c.font      = S.font(bold=True)
            c.fill      = S.fill(S.COL_BG)
            c.alignment = S.align("center", wrap=True)
            if bi == N_BUCKETS - 1:
                c.border = S.right_sep()
        ws.row_dimensions[4].height = 36

        databases        = self._databases_sorted()
        db_bucket_totals = {}
        db_totals        = {}
        grand_buckets    = {b: 0.0 for b in AGING_BUCKETS}
        grand_total      = 0.0
        cur_row          = 5

        for di, db in enumerate(databases):
            db_rows = self.df[self.df["DATABASE"] == db]
            bucket_totals = {
                b: float(db_rows.loc[db_rows["OPEN_GROUP"] == b,
                                     "OPEN_AMOUNT_IN_USD"].sum())
                for b in AGING_BUCKETS
            }
            total_open  = sum(bucket_totals.values())
            pd_total    = sum(bucket_totals[b] for b in AGING_BUCKETS
                              if b.startswith("PD"))
            pct_overdue = pd_total / total_open if total_open else None
            row_bg      = S.ALT_BG if di % 2 == 1 else S.EVEN_BG

            db_bucket_totals[db] = bucket_totals
            db_totals[db]        = total_open

            self._sc(ws, cur_row, COL_DB, db, bold=True, bg=row_bg,
                     halign="left",
                     border=Border(right=S.SEP, bottom=S.THIN_G))

            for bi, bucket in enumerate(AGING_BUCKETS):
                col = COL_BUCKET_START + bi
                val = bucket_totals[bucket]
                brd = Border(right=S.SEP, bottom=S.THIN_G) \
                    if bi == N_BUCKETS - 1 else Border(bottom=S.THIN_G)
                self._sc(ws, cur_row, col,
                         val if val != 0 else None,
                         bg=row_bg, fmt=NUM_USD, border=brd)
                grand_buckets[bucket] += val

            self._sc(ws, cur_row, COL_TOTAL, total_open, bold=True,
                     bg=row_bg, fmt=NUM_USD,
                     border=Border(right=S.SEP, bottom=S.THIN_G))
            self._sc(ws, cur_row, COL_PCT, pct_overdue,
                     bg=row_bg, fmt="0.0%", halign="center",
                     border=Border(right=S.SEP, bottom=S.THIN_G))

            grand_total += total_open
            ws.row_dimensions[cur_row].height = 15
            cur_row += 1

        for col in range(1, TOTAL_COLS + 1):
            cell = ws.cell(row=cur_row - 1, column=col)
            cell.border = Border(top=cell.border.top, left=cell.border.left,
                                 right=cell.border.right, bottom=S.OUTER)

        gt_row = cur_row
        gt_pd  = sum(grand_buckets[b] for b in AGING_BUCKETS if b.startswith("PD"))
        gt_pct = gt_pd / grand_total if grand_total else None

        self._sc(ws, gt_row, COL_DB, "Grand Total", bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, halign="left",
                 border=Border(right=S.SEP))
        for bi, bucket in enumerate(AGING_BUCKETS):
            col = COL_BUCKET_START + bi
            val = grand_buckets[bucket]
            brd = Border(right=S.SEP) if bi == N_BUCKETS - 1 else None
            self._sc(ws, gt_row, col, val if val != 0 else None,
                     bold=True, fg=S.GRAND_FG, bg=S.GRAND_BG,
                     fmt=NUM_USD, border=brd)
        self._sc(ws, gt_row, COL_TOTAL, grand_total, bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, fmt=NUM_USD,
                 border=Border(right=S.SEP))
        self._sc(ws, gt_row, COL_PCT, gt_pct, bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, fmt="0.0%",
                 halign="center", border=Border(right=S.SEP))
        ws.row_dimensions[gt_row].height = 16

        S.apply_outer_border(ws, 1, gt_row, 1, TOTAL_COLS)
        cur_row = gt_row + 1

        # ── SPACER ────────────────────────────────────────────────────────────
        ws.row_dimensions[cur_row].height = 10
        cur_row += 1

        # ── % OF TOTAL TABLE ──────────────────────────────────────────────────

        ws.merge_cells(start_row=cur_row, start_column=COL_BUCKET_START,
                       end_row=cur_row,   end_column=COL_TOTAL - 1)
        c_pct = ws.cell(row=cur_row, column=COL_BUCKET_START,
                        value="Aging Buckets  (% of Total Open)")
        c_pct.font      = S.font(bold=True, color=S.GRP_FG)
        c_pct.fill      = S.fill(S.GRP_BG)
        c_pct.alignment = S.align("center")
        for col in range(COL_BUCKET_START, COL_TOTAL):
            ws.cell(row=cur_row, column=col).fill = S.fill(S.GRP_BG)
        for col in [COL_DB, COL_TOTAL, COL_PCT]:
            ws.cell(row=cur_row, column=col).fill = S.fill(S.TITLE_BG)
        ws.row_dimensions[cur_row].height = 14
        pct_table_start = cur_row
        cur_row += 1

        self._sc(ws, cur_row, COL_DB, "Database", bold=True, bg=S.COL_BG,
                 halign="left", border=S.right_sep())
        self._sc(ws, cur_row, COL_TOTAL, "Total Open\n(USD)", bold=True, bg=S.COL_BG,
                 halign="right", border=S.right_sep(), wrap=True)
        self._sc(ws, cur_row, COL_PCT, "% Past Due", bold=True, bg=S.COL_BG,
                 halign="center", border=S.right_sep(), wrap=True)
        for bi, bucket in enumerate(AGING_BUCKETS):
            col   = COL_BUCKET_START + bi
            label = AGING_LABELS[bucket]
            c     = ws.cell(row=cur_row, column=col, value=label)
            c.font      = S.font(bold=True)
            c.fill      = S.fill(S.COL_BG)
            c.alignment = S.align("center", wrap=True)
            if bi == N_BUCKETS - 1:
                c.border = S.right_sep()
        ws.row_dimensions[cur_row].height = 36
        cur_row += 1

        for di, db in enumerate(databases):
            bucket_totals = db_bucket_totals[db]
            total_open    = db_totals[db]
            pd_total      = sum(bucket_totals[b] for b in AGING_BUCKETS
                                if b.startswith("PD"))
            pct_overdue   = pd_total / total_open if total_open else None
            row_bg        = S.ALT_BG if di % 2 == 1 else S.EVEN_BG

            self._sc(ws, cur_row, COL_DB, db, bold=True, bg=row_bg,
                     halign="left",
                     border=Border(right=S.SEP, bottom=S.THIN_G))

            for bi, bucket in enumerate(AGING_BUCKETS):
                col = COL_BUCKET_START + bi
                val = (bucket_totals[bucket] / total_open) if total_open else None
                brd = Border(right=S.SEP, bottom=S.THIN_G) \
                    if bi == N_BUCKETS - 1 else Border(bottom=S.THIN_G)
                self._sc(ws, cur_row, col,
                         val if val else None,
                         bg=row_bg, fmt="0.0%", halign="center", border=brd)

            self._sc(ws, cur_row, COL_TOTAL, total_open, bold=True,
                     bg=row_bg, fmt=NUM_USD,
                     border=Border(right=S.SEP, bottom=S.THIN_G))
            self._sc(ws, cur_row, COL_PCT, pct_overdue,
                     bg=row_bg, fmt="0.0%", halign="center",
                     border=Border(right=S.SEP, bottom=S.THIN_G))

            ws.row_dimensions[cur_row].height = 15
            cur_row += 1

        for col in range(1, TOTAL_COLS + 1):
            cell = ws.cell(row=cur_row - 1, column=col)
            cell.border = Border(top=cell.border.top, left=cell.border.left,
                                 right=cell.border.right, bottom=S.OUTER)

        gt_row2 = cur_row
        self._sc(ws, gt_row2, COL_DB, "Grand Total", bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, halign="left",
                 border=Border(right=S.SEP))
        for bi, bucket in enumerate(AGING_BUCKETS):
            col = COL_BUCKET_START + bi
            val = (grand_buckets[bucket] / grand_total) if grand_total else None
            brd = Border(right=S.SEP) if bi == N_BUCKETS - 1 else None
            self._sc(ws, gt_row2, col, val if val else None,
                     bold=True, fg=S.GRAND_FG, bg=S.GRAND_BG,
                     fmt="0.0%", halign="center", border=brd)
        self._sc(ws, gt_row2, COL_TOTAL, grand_total, bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, fmt=NUM_USD,
                 border=Border(right=S.SEP))
        self._sc(ws, gt_row2, COL_PCT, gt_pct, bold=True,
                 fg=S.GRAND_FG, bg=S.GRAND_BG, fmt="0.0%",
                 halign="center", border=Border(right=S.SEP))
        ws.row_dimensions[gt_row2].height = 16

        S.apply_outer_border(ws, pct_table_start, gt_row2, 1, TOTAL_COLS)

        ws.column_dimensions["A"].width = 14
        for bi in range(N_BUCKETS):
            ws.column_dimensions[get_column_letter(COL_BUCKET_START + bi)].width = 16
        ws.column_dimensions[get_column_letter(COL_TOTAL)].width = 16
        ws.column_dimensions[get_column_letter(COL_PCT)].width   = 12

        ws.freeze_panes = "A5"
        print(f"  Summary: {len(databases)} databases | Total Open AR ${grand_total:,.0f}")

    # ── DETAIL tab ────────────────────────────────────────────────────────────

    def _write_detail_headers(self, ws: Worksheet, N: int,
                              title: str, subtitle: str) -> None:
        self._title_banner(ws, N, title, subtitle)
        ws.row_dimensions[3].height = 32
        for ci, (hdr, width, _, _, _) in enumerate(DETAIL_COLS, start=1):
            c = ws.cell(row=3, column=ci, value=hdr)
            c.font      = S.font(bold=True)
            c.fill      = S.fill(S.COL_BG)
            c.alignment = S.align("center", wrap=True)
            c.border    = Border(left=S.THIN_G, right=S.THIN_G,
                                 top=S.OUTER,   bottom=S.OUTER)
            ws.column_dimensions[get_column_letter(ci)].width = width

    def _write_detail_rows(self, ws: Worksheet,
                           rows: pd.DataFrame, start_row: int) -> int:
        cur         = start_row
        signal_keys = {"DAYS_OVERDUE", "OPEN_GROUP_LABEL"}

        for row_idx, (_, row_data) in enumerate(rows.iterrows()):
            bucket             = str(row_data.get("OPEN_GROUP", ""))
            sig_color, sig_bold = BUCKET_STYLE.get(bucket, ("000000", False))
            bg_hex             = S.ALT_BG if row_idx % 2 == 1 else S.EVEN_BG

            ws.row_dimensions[cur].height = 15
            for ci, (_, _, data_key, num_fmt, halign) in enumerate(DETAIL_COLS, start=1):
                val = row_data.get(data_key, "")
                if data_key in ("REF_DATE", "DUE_DATE"):
                    val = val.date() if pd.notna(val) and hasattr(val, "date") else None
                elif data_key == "DAYS_OVERDUE":
                    val = int(val) if pd.notna(val) else None

                is_signal = data_key in signal_keys
                c = ws.cell(row=cur, column=ci, value=val)
                c.font          = S.font(bold=sig_bold if is_signal else False,
                                         color=sig_color if is_signal else "000000")
                c.fill          = S.fill(bg_hex)
                c.alignment     = S.align(halign)
                c.number_format = num_fmt
                c.border        = Border(left=S.THIN_G, right=S.THIN_G,
                                         bottom=S.THIN_G)
            cur += 1

        return cur

    def _write_grand_total_row(self, ws: Worksheet, row: int,
                               grand_usd: float, N: int, usd_col: int) -> None:
        ws.row_dimensions[row].height = 16
        for ci in range(1, N + 1):
            c = ws.cell(row=row, column=ci)
            c.fill   = S.fill(S.GRAND_BG)
            c.font   = S.font(bold=True, color=S.GRAND_FG)
            c.border = Border(top=S.OUTER, bottom=S.OUTER,
                              left=S.THIN_G, right=S.THIN_G)
        ws.cell(row=row, column=1, value="GRAND TOTAL").alignment = S.align("left")
        ws.cell(row=row, column=1).font = S.font(bold=True, color=S.GRAND_FG)
        c = ws.cell(row=row, column=usd_col, value=grand_usd)
        c.number_format = "#,##0.00"
        c.alignment     = S.align("right")
        c.font          = S.font(bold=True, color=S.GRAND_FG)
        c.fill          = S.fill(S.GRAND_BG)

    def build_detail_tab(self, df: pd.DataFrame, sheet_name: str,
                         title: str) -> None:
        ws      = self.wb.create_sheet(sheet_name)
        N       = len(DETAIL_COLS)
        usd_col = next(i + 1 for i, (_, _, k, _, _) in enumerate(DETAIL_COLS)
                       if k == "OPEN_AMOUNT_IN_USD")

        self._write_detail_headers(
            ws, N, title,
            f"{len(df):,} open items  |  As of {TODAY.strftime('%B %d, %Y')}",
        )

        cur_row   = 4
        grand_usd = float(df["OPEN_AMOUNT_IN_USD"].sum())
        cur_row   = self._write_detail_rows(ws, df, cur_row)

        self._write_grand_total_row(ws, cur_row, grand_usd, N, usd_col)
        S.apply_outer_border(ws, 1, cur_row, 1, N)
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(N)}3"
        print(f"  {sheet_name}: {len(df):,} rows | Open USD ${grand_usd:,.0f}")

    # ── EXCHANGE RATES tab ────────────────────────────────────────────────────

    def build_exchange_rates_tab(self) -> None:
        ws = self.wb.create_sheet("EXCHANGE RATES")

        TOTAL_COLS = 4
        self._title_banner(
            ws, TOTAL_COLS,
            f"Exchange Rates Used  |  {TODAY.strftime('%B %d, %Y')}",
            "Spot rates applied for USD conversion in this report",
        )

        headers = ["Currency Pair", "From", "To", "Rate (Local per USD)"]
        widths  = [22, 10, 10, 22]
        for ci, (hdr, width) in enumerate(zip(headers, widths), start=1):
            c = ws.cell(row=3, column=ci, value=hdr)
            c.font      = S.font(bold=True)
            c.fill      = S.fill(S.COL_BG)
            c.alignment = S.align("center")
            c.border    = Border(left=S.THIN_G, right=S.THIN_G,
                                 top=S.OUTER,   bottom=S.OUTER)
            ws.column_dimensions[get_column_letter(ci)].width = width
        ws.row_dimensions[3].height = 18

        rate_rows = []
        if "LATEST_X_RATE" in self.df.columns and "ENTERED_CURRENCY" in self.df.columns:
            rate_df = (
                self.df[self.df["ENTERED_CURRENCY"].isin(["RMB", "EURO", "CLP"])]
                [["ENTERED_CURRENCY", "LATEST_X_RATE"]]
                .dropna()
                .drop_duplicates("ENTERED_CURRENCY")
                .sort_values("ENTERED_CURRENCY")
            )
            for _, row in rate_df.iterrows():
                currency      = row["ENTERED_CURRENCY"]
                usd_per_local = float(row["LATEST_X_RATE"])
                local_per_usd = round(1.0 / usd_per_local, 6) if usd_per_local else None
                rate_rows.append((f"{currency} / USD", currency, "USD", local_per_usd))

        rate_rows.append(("USD / USD", "USD", "USD", 1.0))

        for ri, (pair, frm, to, rate_val) in enumerate(rate_rows):
            row_bg = S.ALT_BG if ri % 2 == 1 else S.EVEN_BG
            cur    = 4 + ri
            ws.row_dimensions[cur].height = 15
            self._sc(ws, cur, 1, pair,     bg=row_bg, halign="left",
                     border=Border(left=S.THIN_G, right=S.THIN_G, bottom=S.THIN_G))
            self._sc(ws, cur, 2, frm,      bg=row_bg, halign="center",
                     border=Border(left=S.THIN_G, right=S.THIN_G, bottom=S.THIN_G))
            self._sc(ws, cur, 3, to,       bg=row_bg, halign="center",
                     border=Border(left=S.THIN_G, right=S.THIN_G, bottom=S.THIN_G))
            self._sc(ws, cur, 4, rate_val, bg=row_bg, halign="right",
                     fmt="0.000000",
                     border=Border(left=S.THIN_G, right=S.THIN_G, bottom=S.THIN_G))

        last_row = 4 + len(rate_rows) - 1
        S.apply_outer_border(ws, 1, last_row, 1, TOTAL_COLS)
        print(f"  Exchange Rates: {len(rate_rows)} currencies")

    # ── Orchestrate all tabs ──────────────────────────────────────────────────

    def build(self) -> None:
        self.build_summary_tab()

        for db in self._databases_sorted():
            safe = (db[:31]
                    .replace("/", "-").replace("\\", "-")
                    .replace("*", "").replace("?",  "")
                    .replace("[", "").replace("]",  "")
                    .replace(":", "").upper())
            db_df = self.df[self.df["DATABASE"] == db]
            self.build_detail_tab(
                df         = db_df,
                sheet_name = safe,
                title      = f"Mafco Open AR  |  {db}  |  {TODAY.strftime('%B %d, %Y')}",
            )

        self.build_exchange_rates_tab()


# ═══════════════════════════════════════════════════════════════════════════════
# EMAILER ENTRY POINT
# ═══════════════════════════════════════════════════════════════════════════════

def build_report(df: pd.DataFrame, _orders_df: Any, config: _ReportConfig) -> None:
    """Called by finance_emailer.py. _orders_df is unused (AR has no orders)."""
    wb = Workbook()
    wb.remove(wb.active)
    builder = ARReportBuilder(wb, df)
    builder.build()
    wb.save(config.output_file)


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════

def main() -> None:
    loader = Loader()
    try:
        df = loader.load()
    finally:
        loader.close()

    wb = Workbook()
    wb.remove(wb.active)

    builder = ARReportBuilder(wb, df)
    builder.build()

    wb.save(OUTPUT_FILE)
    print(f"\nSaved → {OUTPUT_FILE}")
    print(f"Tabs:  {[ws.title for ws in wb.worksheets]}")


if __name__ == "__main__":
    main()